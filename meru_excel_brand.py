"""
MERU Excel Brand Configuration
Sourced directly from Excel_Workbook.xlsx template.
Import this module in any MERU Excel generator.

FORMATTING RULES
- Never use merge and center. Always use centerContinuous (Center Across Selection).
- Apply centerContinuous to every cell in the target range, not just the first cell.
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers, NamedStyle
from openpyxl.utils import get_column_letter

# ─── Colors ───────────────────────────────────────────────────────────────────
GREEN      = '09B582'   # MERU primary green (header fill)
CHARCOAL   = '414042'   # body text
DARK_GREEN = '03492E'
NAVY       = '08203C'
WHITE      = 'FFFFFF'
TOTAL_FILL = 'F2F2F2'   # near-white for total rows (theme 0, tint -0.05)

FONT       = 'Segoe UI'

# ─── Number formats (from template) ───────────────────────────────────────────
FMT_NUMBER   = '_(#,##0_);_(\\(#,##0\\)_);" - "'
FMT_CURRENCY = '_("$"#,##0_);_(\\("$"#,##0\\)_);" - "'
FMT_PERCENT  = '_(0.0%_);_(\\(0.0%\\)_);" - "'
FMT_TEXT     = '@'

# ─── Row heights & column widths (from template) ─────────────────────────────
ROW_HEIGHTS = {
    'spacer_top':  5.0,
    'client_name': 25.15,
    'subtitle':    15.0,
    'header':      15.0,
    'data':        15.0,
    'total':       15.0,
    'spacer':      15.0,
}

COL_WIDTHS = {
    'margin_left':  1.82,   # Col A
    'label':        20.18,  # Col B
    'data_narrow':  13.0,
    'data_wide':    15.5,
    'margin_right': 2.63,
}

# ─── Style factories ──────────────────────────────────────────────────────────

def _font(size=11, bold=False, color=None):
    return Font(name=FONT, size=size, bold=bold, color=color or CHARCOAL)

def _fill(hex_color):
    return PatternFill('solid', fgColor=hex_color)

def _border_thin():
    s = Side(style='thin', color='CCCCCC')
    return Border(top=s, bottom=s, left=s, right=s)

def _border_total():
    """Total row: thin top, double bottom."""
    return Border(
        top=Side(style='thin', color='000000'),
        bottom=Side(style='double', color='000000'),
    )

def _align(h='left', v='center', wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


# ─── Cell-level helpers ───────────────────────────────────────────────────────

def style_client_name(cell, value):
    """Client/engagement name header: 16pt bold black, underlined."""
    cell.value = value
    cell.font  = Font(name=FONT, size=16, bold=True, underline='single', color='000000')
    cell.alignment = _align()

def style_subtitle(cell, value):
    """Tab title: italic, black, no underline."""
    cell.value = value
    cell.font  = Font(name=FONT, size=11, italic=True, color='000000')
    cell.alignment = _align()

def style_units(cell, value):
    """Units label: italic, black, single underline."""
    cell.value = value
    cell.font  = Font(name=FONT, size=11, italic=True, underline='single', color='000000')
    cell.alignment = _align()

def style_header(cell, value, align='right'):
    """Column header: green fill, white bold text."""
    cell.value = value
    cell.font  = _font(size=11, bold=True, color=WHITE)
    cell.fill  = _fill(GREEN)
    cell.alignment = _align(h=align)

def style_label_header(cell, value):
    """Leftmost header cell (label column): green fill, white bold, left-aligned."""
    style_header(cell, value, align='left')

def style_section_header(cell, value, col_span_cells=None):
    """Section header spanning multiple cells (like 'Formats' row)."""
    cell.value = value
    cell.font  = _font(size=11, bold=True, color=WHITE)
    cell.fill  = _fill(GREEN)
    cell.alignment = _align(h='centerContinuous')
    if col_span_cells:
        for c in col_span_cells:
            c.fill = _fill(GREEN)
            c.alignment = _align(h='centerContinuous')

def style_data(cell, value=None, fmt=FMT_NUMBER, align='right', bold=False):
    """Standard data cell."""
    if value is not None:
        cell.value = value
    cell.font   = _font(size=11, bold=bold)
    cell.number_format = fmt
    cell.alignment = _align(h=align)

def style_label(cell, value, bold=False, indent=0):
    """Row label (left column)."""
    cell.value = value
    cell.font  = _font(size=11, bold=bold)
    cell.alignment = Alignment(horizontal='left', vertical='center', indent=indent)

def style_total_label(cell, value):
    """Total row label."""
    cell.value = value
    cell.font  = _font(size=11, bold=True)
    cell.fill  = _fill(TOTAL_FILL)
    cell.border = _border_total()
    cell.alignment = _align(h='left')

def style_total_data(cell, value=None, fmt=FMT_CURRENCY):
    """Total row data cell."""
    if value is not None:
        cell.value = value
    cell.font   = _font(size=11, bold=True)
    cell.fill   = _fill(TOTAL_FILL)
    cell.border = _border_total()
    cell.number_format = fmt
    cell.alignment = _align(h='right')


# ─── Sheet-level setup ────────────────────────────────────────────────────────

def setup_workbook(wb):
    """
    Set Segoe UI as the workbook-level default font.
    Call this immediately after creating a Workbook() instance, before adding sheets.
    This prevents Excel from falling back to Arial when a cell is edited (e.g. typing
    a decimal via hotkey), which happens when no workbook-level default is declared.
    """
    wb.properties.creator = 'MERU'
    # openpyxl stores fonts in an IndexedList; index 0 is the workbook default.
    # Replace it in place (do not reassign the list — that breaks the .add() method).
    wb._fonts[0] = Font(name=FONT, size=11, color=CHARCOAL)


def setup_sheet(ws, tab_color=None):
    """Apply standard column widths and freeze panes."""
    ws.column_dimensions['A'].width = COL_WIDTHS['margin_left']
    ws.column_dimensions['B'].width = COL_WIDTHS['label']
    # Data columns C–F
    for col in ['C', 'D', 'E', 'F']:
        ws.column_dimensions[col].width = COL_WIDTHS['data_narrow']
    ws.column_dimensions['G'].width = COL_WIDTHS['margin_right']

    ws.row_dimensions[1].height = ROW_HEIGHTS['spacer_top']
    ws.row_dimensions[2].height = ROW_HEIGHTS['client_name']

    ws.freeze_panes = 'C8'   # freeze rows 1-7 and col A-B

    if tab_color:
        ws.sheet_properties.tabColor = tab_color

def apply_print_settings(ws, orientation='landscape'):
    from openpyxl.worksheet.page import PageMargins
    ws.page_setup.orientation = orientation
    ws.page_setup.fitToPage   = True
    ws.page_setup.fitToWidth  = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.75, bottom=0.75)


# ─── Cover sheet ──────────────────────────────────────────────────────────────

COVER_COLS = [chr(c) for c in range(ord('A'), ord('O') + 1)]  # A through O

def build_cover_sheet(ws, client_name, workbook_type, date_label=None):
    """
    Build a MERU-branded cover sheet matching Excel_Workbook.xlsx template.

    Layout
    ------
    Rows  1-4  : Top green band (#09B582), full width (A-O)
    Rows  5-9  : White space
    Row   10   : Client name  -- 24pt bold, centerContinuous, row height 35
    Rows 11-12 : White space
    Row   13   : Workbook type -- 18pt bold, centerContinuous, row height 26
    Row   14   : Small spacer (height 5.25)
    Row   15   : Date / subtitle (optional) -- 14pt, centerContinuous, height 21
    Rows 16-30 : White space
    Rows 31-34 : Bottom green band (#09B582), full width (A-O)

    Parameters
    ----------
    client_name   : str  e.g. 'Cinterra Energy'
    workbook_type : str  e.g. 'Financial Model'
    date_label    : str  optional, e.g. 'June 2026 | Confidential'
    """
    green_fill = PatternFill('solid', fgColor=GREEN)
    white_fill = PatternFill('solid', fgColor=WHITE)

    # Column widths
    cover_col_widths = {
        'A': 1.63, 'B': 8.27, 'C': 8.27, 'D': 8.27, 'E': 8.27,
        'F': 8.27, 'G': 8.27, 'H': 8.27, 'I': 8.27, 'J': 8.27,
        'K': 8.27, 'L': 8.27, 'M': 11.73, 'N': 9.91, 'O': 8.27,
    }
    for col, w in cover_col_widths.items():
        ws.column_dimensions[col].width = w

    # Green bands: top rows 1-4, bottom rows 31-34
    for band_rows in [range(1, 5), range(31, 35)]:
        for r in band_rows:
            ws.row_dimensions[r].height = 15.0
            for col in COVER_COLS:
                ws[f'{col}{r}'].fill = green_fill
    ws.row_dimensions[1].height = 5.25  # narrow top strip

    # White middle rows 5-30
    for r in range(5, 31):
        for col in COVER_COLS:
            ws[f'{col}{r}'].fill = white_fill

    center_across = Alignment(horizontal='centerContinuous', vertical='center')

    def _center_row(row, value, font, row_height):
        """Set value/font on A, then apply centerContinuous to all cols A-O."""
        ws.row_dimensions[row].height = row_height
        ws[f'A{row}'].value = value
        ws[f'A{row}'].font  = font
        for col in COVER_COLS:
            ws[f'{col}{row}'].alignment = center_across

    # Client name -- row 10
    _center_row(10, client_name,   Font(name=FONT, size=24, bold=True,  color=CHARCOAL), 35.0)

    # Workbook type -- row 13
    _center_row(13, workbook_type, Font(name=FONT, size=18, bold=True,  color=CHARCOAL), 26.0)

    # Date / subtitle -- row 15 (optional)
    ws.row_dimensions[14].height = 5.25
    ws.row_dimensions[15].height = 21.0
    if date_label:
        _center_row(15, date_label, Font(name=FONT, size=14, bold=False, color='888888'), 21.0)

    # Tab styling
    ws.sheet_properties.tabColor = GREEN
    ws.sheet_view.showGridLines  = False


# ─── Post-save font patch ──────────────────────────────────────────────────────

def finalize_workbook(filepath):
    """
    Post-processing step: replace all non-Segoe UI fonts in styles.xml with
    Segoe UI 11pt. Call this after wb.save(filepath).

    Why this is needed: openpyxl hard-codes Arial 10pt fonts for its built-in
    named styles (Comma, Currency, Percent). When Excel applies a number format
    via hotkey (e.g. comma, decimal, currency), it references those named styles
    and their Arial font. This patch replaces every Arial instance in the file so
    Segoe UI is preserved regardless of how the cell is formatted.
    """
    import zipfile, shutil, os
    import xml.etree.ElementTree as ET

    NS   = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    ET.register_namespace('', NS)

    tmp = filepath + '.tmp'

    with zipfile.ZipFile(filepath, 'r') as zin, \
         zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:

        for item in zin.infolist():
            data = zin.read(item.filename)

            if item.filename == 'xl/styles.xml':
                root = ET.fromstring(data)
                fonts_el = root.find(f'{{{NS}}}fonts')
                if fonts_el is not None:
                    for font_el in fonts_el.findall(f'{{{NS}}}font'):
                        name_el = font_el.find(f'{{{NS}}}name')
                        sz_el   = font_el.find(f'{{{NS}}}sz')
                        if name_el is not None and name_el.get('val', '').lower() != 'segoe ui':
                            name_el.set('val', FONT)
                        if sz_el is not None and sz_el.get('val') == '10':
                            sz_el.set('val', '11')
                data = ET.tostring(root, encoding='unicode', xml_declaration=False).encode('utf-8')
                data = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + data

            zout.writestr(item, data)

    os.replace(tmp, filepath)
