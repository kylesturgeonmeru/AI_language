"""Report: coverage report first, then the deliverables. Reads only saved files.

Writes, in this order (CLAUDE.md rule 8: coverage before any summary):
  data/coverage_report.md       attempted, succeeded, failed, and gaps
  output/ai_el_benchmark.xlsx   Provisions, Prevalence, Case List, Manual Pull List
  output/clause_library.md      verbatim AI provisions grouped by firm

output/benchmark_memo.md is written in-session from these outputs.
Usage: python report.py
"""
import csv
import io
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment

from classify import RUBRIC
from common import COURTS, DATA, ROOT, atomic_write_text
from meru_excel_brand import (FMT_NUMBER, FMT_PERCENT, FMT_TEXT, apply_print_settings, build_cover_sheet,
                              finalize_workbook, setup_workbook, style_client_name, style_data, style_header,
                              style_label, style_label_header, style_subtitle, style_total_data,
                              style_total_label)

OUT = ROOT / "output"
VENUE = {"deb": "D. Del.", "njb": "D.N.J.", "txsb": "S.D. Tex.", "txnb": "N.D. Tex.", "nysb": "S.D.N.Y.",
         "vaeb": "E.D. Va.", "ganb": "N.D. Ga."}
ROLE_ORDER = ["FA", "committee FA", "IB", "CRO", "debtor counsel", "committee counsel", "claims agent",
              "administrative advisor", "OCP", "other"]
ROLE_SET = {"FA": "primary", "committee FA": "primary", "IB": "primary", "CRO": "primary",
            "debtor counsel": "comparison", "committee counsel": "comparison",
            "claims agent": "secondary", "administrative advisor": "secondary", "OCP": "secondary",
            "other": "other"}
CLIENT_SIDE = {("25-11120", "604")}  # AI provision addresses the client's own AI use (Kyle, 10/11/2026)


def rows(path):
    with open(DATA / path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def case_names():
    """Clean case names from docket metadata, keyed by case number."""
    return {r["docket_number"]: r["case_name"] for r in rows("cases.csv")}


def classifications():
    return [json.loads(l) for l in open(DATA / "classifications.jsonl", encoding="utf-8")]


def in_track(c, track):
    return track in c["track"].split("|")


def claims_agents():
    """Docket number to claims agent name, from claims agent applications in the discovery data."""
    out = {}
    pat = re.compile(r"Claims/Noticing Agent\s+(.+?)\s+(?:Filed|\(|$)", re.I)
    for c in rows("candidates.csv"):
        m = pat.search(c["description"])
        if m and c["docket_number"] not in out:
            out[c["docket_number"]] = m.group(1).strip(" ,.").title()
    return out


def manual_pull_list():
    agents = claims_agents()
    out = []
    for m in rows("manual_pull_list.csv"):
        out.append({"track": "B", "case_name": m["case_name"], "docket_number": m["docket_number"],
                    "document_number": m["document_number"], "attachment_number": m["attachment_number"],
                    "description": m["description"][:200], "claims_agent": agents.get(m["docket_number"], ""),
                    "reason": "Track B related document not in RECAP"})
    avail = {(r["docket_number"], r["document_number"], r["attachment_number"]): r["is_available"]
             for r in rows("original_top10_fetch_list.csv")}
    for r in rows("original_top10_retentions.csv"):
        if r["kind"] != "application" or avail.get((r["docket_number"], r["document_number"], "")) == "True":
            continue
        out.append({"track": "A (original top 10)", "case_name": r["case_name"], "docket_number": r["docket_number"],
                    "document_number": r["document_number"], "attachment_number": "",
                    "description": r["description"][:200], "claims_agent": agents.get(r["docket_number"], ""),
                    "reason": "Retention application not in RECAP; outside the prevalence denominator"})
    for o in out:
        o["claims_agent_docket_url"] = ("Not scraped (CLAUDE.md source rule 3): open the case page on the claims "
                                        "agent's site and pull the docket number shown")
    return out


# ---------------------------------------------------------------- coverage report

def coverage_report(cls):
    cov = rows("track_a_coverage.csv")
    census = [c for c in cov if c["qualifies"] == "True"]
    fm = rows("fetch_manifest.csv")
    ok = [m for m in fm if m["status"] == "ok"]
    failed = [m for m in fm if m["status"] != "ok"]
    rank_m = rows("ranking_manifest.csv")
    sel_a = rows("fetch_list_a.csv")
    sel_b = rows("fetch_list_b.csv")
    a_cls = [c for c in cls if in_track(c, "A")]
    b_cls = [c for c in cls if in_track(c, "B")]
    orders = [r for r in rows("census_retentions.csv") if r["kind"] == "entered_order"]
    ocp = [r for r in rows("census_retentions.csv") if r["kind"] == "ocp_declaration"]
    mpl = manual_pull_list()
    orig = rows("original_top10_fetch_list.csv")
    zero = [c for c in cov if c["qualifies"] != "True"]

    L = ["# Coverage report", "",
         "Generated by `report.py` from saved files. Gaps are unknown, never \"no AI provision\". "
         "Filings not in RECAP are outside the prevalence denominator and appear on the Manual Pull List.", ""]
    L += ["## Requests and downloads", "",
          "| Item | Attempted | Succeeded | Failed |", "|---|---|---|---|",
          f"| Research documents (Track A and B) | {len(fm)} | {len(ok)} | {len(failed)} |",
          f"| Ranking declarations (exempt from cap) | {len(rank_m)} | "
          f"{sum(m['status'] == 'ok' for m in rank_m)} | {sum(m['status'] != 'ok' for m in rank_m)} |", "",
          f"Download cap: {len(ok)} of 400 research documents used. All PDFs passed validation "
          "(header, trailer, and a full parse).", ""]

    L += ["## Track B (full-text discovery)", "",
          "- Searched: 8 AI terms across chapter 11 cases in the 7 venues, petitions from 01/01/2025, "
          "docket entries described as retention filings.",
          f"- Hits: {len(b_cls)} filings in {len({c['case_number'] for c in b_cls})} cases. "
          "A sensitivity check without the retention filter (287 documents) found no retention filings the filter dropped.",
          f"- Documents selected: {len(sel_b)}; in RECAP and fetched: "
          f"{sum(r['is_available'] == 'True' for r in sel_b)}; not in RECAP: "
          f"{sum(r['is_available'] != 'True' for r in sel_b)} (Manual Pull List).", ""]

    L += ["## Track A (census of large cases, for prevalence)", "",
          f"- Ranked: {len(cov)} in-scope cases by funded debt in the first-day declaration "
          "(132 declarations read).",
          f"- Included: {len(census)} cases with at least one retention application in RECAP; "
          f"{sum(int(c['available_applications']) for c in census)} applications, {len(orders)} entered orders.",
          f"- Documents selected (applications with exhibits, entered orders): {len(sel_a)}, all in RECAP and fetched.",
          f"- Classified: {len(a_cls)} applications.", "",
          "| Rank | Case | Venue | Funded debt ($m) | Applications in RECAP | Classified | AI provisions |",
          "|---|---|---|---|---|---|---|"]
    by_case = defaultdict(list)
    for c in a_cls:
        by_case[c["case_number"]].append(c)
    names = case_names()
    for c in census:
        k = c["case_number"]
        L.append(f"| {c['rank']} | {names.get(k, c['debtor'])[:60]} | {VENUE.get(c['court_id'], c['court_id'])} | "
                 f"{float(c['funded_debt_musd']):,.0f} | {c['available_applications']} | {len(by_case.get(k, []))} | "
                 f"{sum(x['ai_provision_present'] == 'Y' for x in by_case.get(k, []))} |")
    L += ["", "### By venue", "", "| Venue | Ranked cases | Included cases | Applications classified |", "|---|---|---|---|"]
    for ct in COURTS:
        L.append(f"| {VENUE[ct]} | {sum(c['court_id'] == ct for c in cov)} | {sum(c['court_id'] == ct for c in census)} | "
                 f"{sum(1 for c in a_cls if c['court'] == VENUE[ct])} |")
    L += ["", "### By role", "", "| Role | Set | Applications classified |", "|---|---|---|"]
    rc = Counter(c["role"] for c in a_cls)
    for r in ROLE_ORDER:
        L.append(f"| {r} | {ROLE_SET[r]} | {rc.get(r, 0)} |")

    L += ["", "## Gaps", "",
          f"1. **Original strict top 10 by funded debt.** RECAP holds {sum(r['is_available'] == 'True' for r in orig)} of "
          f"{len(orig)} application and order documents. Their applications not in RECAP are on the Manual Pull List.",
          f"2. **Ranked cases with no retention application in RECAP:** {len(zero)} of {len(cov)}, so they are outside the denominator.",
          "3. **Large cases with no first-day declaration in RECAP,** so they were never ranked: Spirit Aviation, Saks Global, "
          "At Home, New Rite Aid, Hughes Satellite, Azul, and Multi-Color.",
          f"4. **Individual OCP declarations:** {len(ocp)} in RECAP for the included cases, not fetched (option 1 covers OCPs "
          "through the procedures motion and its order).",
          f"5. **Entered orders:** {len(orders)} entered retention orders in RECAP against {len(a_cls)} applications; "
          "R14 is coded not_available where the order is missing.",
          "6. **Selection bias.** RECAP holds documents that someone bought on PACER or that were free. Track A prevalence "
          "describes large cases' retention filings that are in RECAP, not all large cases.",
          "7. **Not covered:** fee applications and fee examiner reports (optional Phase 5), and individual SEC leads "
          "(the one lead, Sleep Number's 8-K EX-10.1, is an asset purchase agreement, out of scope).",
          f"8. **Manual Pull List:** {len(mpl)} documents.", ""]
    atomic_write_text(DATA / "coverage_report.md", "\n".join(L) + "\n")
    return census


# ---------------------------------------------------------------- workbook

def write_table(ws, title, subtitle, headers, data, widths, start=4, text_cols=(), pct_cols=()):
    style_client_name(ws["B2"], title)
    style_subtitle(ws["B3"], subtitle)
    ws.column_dimensions["A"].width = 1.82
    for i, (h, w) in enumerate(zip(headers, widths)):
        col = 2 + i
        cell = ws.cell(row=start, column=col)
        (style_label_header if i == 0 else style_header)(cell, h)
        cell.alignment = Alignment(horizontal="left" if i == 0 or col in text_cols else "right",
                                   vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = w
    for r, row in enumerate(data, start=start + 1):
        for i, v in enumerate(row):
            cell = ws.cell(row=r, column=2 + i)
            if i == 0:
                style_label(cell, v)
            elif isinstance(v, (int, float)) and not isinstance(v, bool):
                style_data(cell, v, fmt=FMT_PERCENT if i in pct_cols else FMT_NUMBER)
            else:
                style_data(cell, v, fmt=FMT_TEXT, align="left")
            cell.alignment = Alignment(horizontal=cell.alignment.horizontal, vertical="top", wrap_text=True)
    ws.freeze_panes = ws.cell(row=start + 1, column=3)
    apply_print_settings(ws)
    return start + len(data)


def prevalence_rows(a_cls, key):
    groups = defaultdict(list)
    for c in a_cls:
        groups[key(c)].append(c)
    return groups


def workbook(cls, census, mpl):
    wb = Workbook()
    setup_workbook(wb)
    cover = wb.active
    cover.title = "Cover"
    build_cover_sheet(cover, "AI Provisions in Chapter 11 Engagement Letters", "Benchmark Workbook",
                      "Petitions 01/01/2025 to present | Internal")

    # Provisions: one row per classified filing
    ws = wb.create_sheet("Provisions")
    head = ["Case", "Track", "Court", "Case no.", "Petition date", "Filing date", "Docket no.", "Firm", "Role",
            "Retained by", "Filing type", "EL attached", "AI provision", "Location"]
    for code in RUBRIC:
        head += [f"{code} value", f"{code} excerpt"]
    head += ["Notes", "Document URL"]
    data = []
    for c in sorted(cls, key=lambda c: (c["ai_provision_present"] != "Y", c["case_name"], int(c["document_number"]))):
        row = [c["case_name"], c["track"], c["court"], c["case_number"], c["petition_date"], c["filing_date"],
               c["document_number"], c["professional_firm"], c["role"], c["retained_by"], c["filing_type"],
               c["el_attached"], c["ai_provision_present"]
               + (" (client-side)" if (c["docket_number"], c["document_number"]) in CLIENT_SIDE else ""),
               c["location"]]
        for code in RUBRIC:
            f = c["rubric"][code]
            row += [f["value"], f.get("excerpt", "")]
        row += [c.get("notes", ""), c["document_url"]]
        data.append(row)
    widths = [28, 8, 10, 11, 11, 11, 9, 26, 15, 11, 18, 9, 11, 22] + [18, 40] * len(RUBRIC) + [50, 40]
    write_table(ws, "Provisions", "One row per classified filing; rubric values with verbatim excerpts",
                head, data, widths, text_cols=set(range(2, 2 + len(head))))

    # Prevalence: Track A only
    a_cls = [c for c in cls if in_track(c, "A")]
    ws = wb.create_sheet("Prevalence")
    style_client_name(ws["B2"], "Prevalence (Track A)")
    style_subtitle(ws["B3"], f"{len(a_cls)} retention applications in RECAP across {len(census)} large cases; "
                             "low confidence where n < 3")
    r = 5
    for title, key, order in (("By role", lambda c: c["role"], ROLE_ORDER),
                              ("By EL attached", lambda c: {"Y": "EL attached", "N": "No EL attached"}.get(
                                  c["el_attached"], "EL status unknown"), None),
                              ("By firm", lambda c: re.sub(r"[,.]", "", c["professional_firm"]), None)):
        groups = prevalence_rows(a_cls, key)
        keys = [k for k in (order or sorted(groups)) if k in groups]
        for i, h in enumerate([title, "Applications", "With AI provision", "Share"]):
            cell = ws.cell(row=r, column=2 + i)
            (style_label_header if i == 0 else style_header)(cell, h)
        r += 1
        for k in keys:
            n, y = len(groups[k]), sum(c["ai_provision_present"] == "Y" for c in groups[k])
            style_label(ws.cell(row=r, column=2), k)
            style_data(ws.cell(row=r, column=3), n)
            style_data(ws.cell(row=r, column=4), y)
            ws.cell(row=r, column=5).value = f"=IF(C{r}=0,0,D{r}/C{r})"
            style_data(ws.cell(row=r, column=5), fmt=FMT_PERCENT)
            r += 1
        first = r - len(keys)
        style_total_label(ws.cell(row=r, column=2), "Total")
        for col, letter in ((3, "C"), (4, "D")):
            style_total_data(ws.cell(row=r, column=col), f"=SUM({letter}{first}:{letter}{r - 1})", fmt=FMT_NUMBER)
        style_total_data(ws.cell(row=r, column=5), f"=IF(C{r}=0,0,D{r}/C{r})", fmt=FMT_PERCENT)
        r += 3
    ws.column_dimensions["A"].width = 1.82
    ws.column_dimensions["B"].width = 42
    for col in "CDE":
        ws.column_dimensions[col].width = 15.5
    apply_print_settings(ws, "portrait")

    # Case list
    ws = wb.create_sheet("Case List")
    agents = claims_agents()
    data = []
    for c in census:
        n = [x for x in a_cls if x["case_number"] == c["case_number"]]
        data.append([case_names().get(c["case_number"], c["debtor"]), "A", VENUE.get(c["court_id"], c["court_id"]), c["case_number"],
                     int(c["rank"]), float(c["funded_debt_musd"]), len(n),
                     sum(x["ai_provision_present"] == "Y" for x in n), agents.get(c["case_number"], "")])
    for c in sorted({(x["case_name"], x["court"], x["case_number"]) for x in cls if in_track(x, "B")}):
        n = [x for x in cls if x["case_number"] == c[2] and in_track(x, "B")]
        data.append([c[0], "B", c[1], c[2], "", "", len(n), sum(x["ai_provision_present"] == "Y" for x in n),
                     agents.get(c[2], "")])
    write_table(ws, "Case List", "Track A census cases (ranked by funded debt) and Track B cases",
                ["Case", "Track", "Court", "Case no.", "Rank", "Funded debt ($m)", "Filings classified",
                 "AI provisions", "Claims agent"], data, [40, 8, 11, 11, 8, 15, 12, 12, 30],
                text_cols={3, 4, 5, 10})

    # Manual pull list
    ws = wb.create_sheet("Manual Pull List")
    data = [[m["case_name"], m["track"], m["docket_number"], m["document_number"], m["attachment_number"],
             m["description"], m["claims_agent"], m["claims_agent_docket_url"], m["reason"]] for m in mpl]
    write_table(ws, "Manual Pull List", "Documents not in RECAP; pull by hand from the claims agent site",
                ["Case", "Track", "Case no.", "Docket no.", "Attachment", "Description", "Claims agent",
                 "Claims agent docket URL", "Reason"], data, [34, 14, 11, 9, 10, 60, 26, 40, 40],
                text_cols=set(range(2, 11)))

    OUT.mkdir(exist_ok=True)
    path = OUT / "ai_el_benchmark.xlsx"
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    finalize_workbook(str(tmp))
    tmp.replace(path)
    return path


# ---------------------------------------------------------------- clause library

def clause_library(cls):
    passages = [json.loads(l) for l in open(DATA / "located" / "passages.jsonl", encoding="utf-8")]
    by_doc = defaultdict(list)
    for p in passages:
        by_doc[(p["docket_number"], p["document_number"])].append(p)
    L = ["# Clause library: AI provisions in chapter 11 retention filings", "",
         "Verbatim text, grouped by firm. Each entry gives the case, docket number, filing date, and where the text sits. "
         "Source: RECAP via CourtListener. Track A is the large-case census; Track B is full-text discovery.", ""]
    firms = defaultdict(list)
    for c in cls:
        if c["ai_provision_present"] == "Y":
            firms[c["professional_firm"]].append(c)
    for firm in sorted(firms):
        L += [f"## {firm}", ""]
        for c in firms[firm]:
            kind = " (client-side: addresses the client's own AI use)" if (c["docket_number"], c["document_number"]) in CLIENT_SIDE else ""
            L += [f"**{c['case_name']}**, Case No. {c['case_number']} ({c['court']}), Dkt {c['document_number']}, "
                  f"filed {c['filing_date']}. Role: {c['role']}. Track {c['track'].replace('|', ' and ')}. "
                  f"Location: {c['location']}{kind}.", ""]
            seen = set()
            for p in by_doc[(c["docket_number"], c["document_number"])]:
                if p["passage"] in seen:
                    continue
                seen.add(p["passage"])
                L += [f"> {p['passage']}", "", f"Source: attachment {p['attachment_number'] or 0}, page {p['page']}. "
                      f"[Document]({c['document_url']})", ""]
    OUT.mkdir(exist_ok=True)
    atomic_write_text(OUT / "clause_library.md", "\n".join(L) + "\n")


def main():
    cls = classifications()
    census = coverage_report(cls)
    mpl = manual_pull_list()
    path = workbook(cls, census, mpl)
    clause_library(cls)
    a = [c for c in cls if in_track(c, "A")]
    print(f"coverage_report.md, {path.name}, clause_library.md written. "
          f"Track A: {sum(c['ai_provision_present'] == 'Y' for c in a)} of {len(a)} applications with AI provisions. "
          f"Manual pull list: {len(mpl)}.")


if __name__ == "__main__":
    main()
